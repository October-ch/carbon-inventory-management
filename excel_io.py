"""Generated-only XLSX input. Hash allowlist is checked before parsing a file."""
import hashlib
import io
import json
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from inventory import ROOT, FIELDS, SHEETS, build_payload, dec, emission, load_rows, validate_rows

INPUT_DIR = ROOT / "data" / "可上傳範例"
REGISTRY = INPUT_DIR / "允許匯入清單.json"
HEADERS = ["底稿列ID", "工作表", "期間", "項目", "ISO類別", "GHG分類", "活動量", "單位", "排放係數", "係數單位", "憑證編號", "數據來源", "係數版本", "排放量(tCO₂e)"]


def sha256(content):
    return hashlib.sha256(content).hexdigest()


def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="235C4A")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for column in ws.columns:
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = min(45, max(16, max(len(str(c.value or "")) for c in column) * 1.35))


def workbook_bytes(rows, version):
    wb = Workbook()
    info = wb.active
    info.title = "使用說明"
    for record in [
        ["項目", "內容"], ["檔案身分", "CARBON-DEMO-114"], ["版本", version],
        ["組織", "澄源科技股份有限公司（虛構）"], ["期間", "114年1月至12月"],
        ["數據說明", "活動量、係數及憑證均自行模擬，不可用於正式盤查。"],
        ["計算方式", "Python計算活動量×kgCO₂e係數÷1000；Excel內無公式、巨集或外部連線。"],
        ["可上傳範圍", "本原型僅接受系統產生之原檔，修改或重新儲存後不接受；可在介面修改活動量。"],
    ]:
        info.append(record)
    style_sheet(info)
    info.column_dimensions["B"].width = 100
    for sheet in SHEETS:
        ws = wb.create_sheet(sheet)
        ws.append(HEADERS)
        for row in [r for r in rows if r["sheet"] == sheet]:
            values = [row[field] for field in FIELDS]
            for index in [6, 8]:
                values[index] = float(values[index])
            ws.append(values + [float(emission(row))])
        for cells in ws.iter_rows(min_row=2):
            for index in [6, 8, 13]:
                cells[index].number_format = "#,##0.000000"
        style_sheet(ws)
    mapping = wb.create_sheet("報告欄位對照")
    mapping.append(["報告欄位ID", "欄位意義", "計算結果", "單位", "來源底稿列ID"])
    for key, metric in build_payload(rows)["metrics"].items():
        if key in ["total", "cat:1", "cat:2", "cat:3"] or key.startswith("sum:"):
            mapping.append([key, metric["label"], float(metric["raw_value"]), metric["unit"], ", ".join(metric["rows"])])
    style_sheet(mapping)
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def ensure_examples():
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    if REGISTRY.exists():
        return json.loads(REGISTRY.read_text(encoding="utf-8"))
    base = load_rows()
    changed = [dict(row) for row in base]
    changed[0]["activity"] = "69000"
    changed[11]["activity"] = "70800"
    entries = []
    for version, rows, description in [
        ("v1", base, "初始底稿：總量 447.141 tCO₂e"),
        ("v2", changed, "更新底稿：1月用電增加1,000度、12月增加2,000度；總量448.551 tCO₂e"),
    ]:
        filename = f"114年_模擬企業盤查底稿_{version}.xlsx"
        content = workbook_bytes(rows, version)
        (INPUT_DIR / filename).write_bytes(content)
        entries.append({"filename": filename, "version": version, "sha256": sha256(content), "description": description})
    REGISTRY.write_text(json.dumps(entries, ensure_ascii=False, indent=2), encoding="utf-8")
    return entries


def import_generated_excel(content):
    if not content or len(content) > 500_000:
        raise ValueError("請上傳 500 KB 以內的系統模擬 Excel 原檔")
    entry = next((e for e in ensure_examples() if e["sha256"] == sha256(content)), None)
    if entry is None:
        raise ValueError("此檔案不在本專案產生的允許清單內。僅接受下載的 v1／v2 Excel 原檔；修改或重新儲存的檔案亦不接受。")
    # Do not parse unrecognized files, evaluate formulas, or retain external links.
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    rows = []
    try:
        for sheet in SHEETS:
            ws = wb[sheet]
            if ws.max_row > 201 or ws.max_column != len(HEADERS):
                raise ValueError("底稿列數或欄位格式不符")
            records = ws.iter_rows()
            if [c.value for c in next(records)] != HEADERS:
                raise ValueError("底稿欄位標題不符")
            for excel_row, cells in enumerate(records, 2):
                if any(c.data_type == "f" for c in cells):
                    raise ValueError("本原型不接受公式儲存格")
                row = {key: str(cell.value) if cell.value is not None else "" for key, cell in zip(FIELDS, cells)}
                if row["sheet"] != sheet:
                    raise ValueError("工作表與資料列標記不一致")
                if abs(emission(row) - dec(cells[-1].value)) > Decimal("0.000000001"):
                    raise ValueError("底稿計算結果與系統重算不符")
                row.update({"source_book": entry["filename"], "source_sheet": sheet, "source_excel_row": excel_row})
                rows.append(row)
        validate_rows(rows)
        if len(rows) != 34:
            raise ValueError("模擬底稿應包含34筆活動資料")
    finally:
        wb.close()
    return rows, dict(entry)
